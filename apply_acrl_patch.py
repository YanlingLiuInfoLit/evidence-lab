import json,re,shutil
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

root=Path('/mnt/data/acrl_patch_work')
# Cases
p=root/'data/cases.public.json'
d=json.load(open(p,encoding='utf-8'))
for c in d['cases']:
    if c['bank_number']==18:
        c['transfer_principle']='Match health claims to the population, duration, outcomes, independence of the evidence, and financial relationships surrounding the recommendation.'
    if c['bank_number']==21:
        c['transfer_principle']='Money and personal information both have value. Verify unexpected opportunities independently before either changes hands.'
p.write_text(json.dumps(d,ensure_ascii=False,indent=2),encoding='utf-8')
# Embedded BANK in play
playp=root/'play.html'; s=playp.read_text(encoding='utf-8')
newbank=json.dumps(d['cases'],ensure_ascii=False,separators=(',',':'))
m=re.search(r'const BANK=(\[.*?\]),\s*P=',s,re.S)
assert m
s=s[:m.start(1)]+newbank+s[m.end(1):]
playp.write_text(s,encoding='utf-8')

# Public docs small additions
readme=root/'README.md'
r=readme.read_text(encoding='utf-8')
if 'ACRL Framework' not in r:
    r += '\n\n## Pedagogical alignment\n\nEvidence Lab is informed by and aligned with the *ACRL Framework for Information Literacy for Higher Education* (2026 revision). It is not intended as a comprehensive implementation or assessment of every Framework knowledge practice or disposition. The strongest alignments are with Authority Is Constructed and Contextual, Information Creation as a Process, Research as Inquiry, Scholarship as Conversation, and Searching as Strategic Exploration; Information Has Value is addressed selectively through cases involving sponsorship, financial incentives, access, personal information, and related decisions.\n'
    readme.write_text(r,encoding='utf-8')

edu=root/'educators.html'
e=edu.read_text(encoding='utf-8')
if 'ACRL Framework' not in e:
    marker='<h2>Teaching resources</h2>'
    insert='<div class="card"><h2>ACRL Framework alignment</h2><p>Evidence Lab is informed by and aligned with the <em>ACRL Framework for Information Literacy for Higher Education</em> (2026 revision). It is not intended as a comprehensive assessment of every Framework knowledge practice or disposition. See the Teaching Guide for the six-frame mapping.</p></div>'
    if marker in e: e=e.replace(marker,insert+marker,1)
    else: e=e.replace('</main>',insert+'</main>',1)
    edu.write_text(e,encoding='utf-8')

# Teaching Guide DOCX: insert section before current section 9.
docp=root/'educators/Evidence_Lab_Teaching_Guide_v1.0.docx'
doc=Document(docp)
# Renumber existing headings 9-14 -> 10-15
renames={
'9. Accessibility and multimodal evidence':'10. Accessibility and multimodal evidence',
'10. Privacy and assessment':'11. Privacy and assessment',
'11. Reuse and localization':'12. Reuse and localization',
'12. Spoiler-free 36-case catalog':'13. Spoiler-free 36-case catalog',
'13. Public vs. confidential educator materials':'14. Public vs. confidential educator materials',
'14. Citation':'15. Citation'
}
for par in doc.paragraphs:
    if par.text in renames: par.text=renames[par.text]; par.style='Heading 1'
# locate insertion point
anchor=next(par for par in doc.paragraphs if par.text=='10. Accessibility and multimodal evidence')

def insert_before(anchor,text,style=None):
    new_p=OxmlElement('w:p')
    anchor._p.addprevious(new_p)
    p=Paragraph(new_p,anchor._parent)
    if style: p.style=style
    p.add_run(text)
    return p
insert_before(anchor,'9. Alignment with the ACRL Framework (2026 revision)','Heading 1')
insert_before(anchor,'Evidence Lab is informed by and aligned with the ACRL Framework for Information Literacy for Higher Education, including revisions approved September 2, 2026. The game is not intended as a comprehensive implementation or assessment of every Framework knowledge practice or disposition. Coverage is intentionally uneven because cases are designed around authentic information problems rather than a six-frame quota.','Normal')
for text in [
'Authority Is Constructed and Contextual — very strong: students evaluate expertise, credibility, incentives, institutional authority, and AI authority in context.',
'Information Creation as a Process — very strong: students trace how posts, images, clips, translations, summaries, citations, and other information products are created, transformed, and circulated.',
'Information Has Value — selective but meaningful: sponsorship, financial incentives, access barriers, money, personal information, attribution, and information exchange appear in relevant cases.',
'Research as Inquiry — very strong: students investigate uncertainty, identify gaps, gather evidence, update conclusions, and sometimes conclude that available evidence is insufficient.',
'Scholarship as Conversation — strong: students trace claims and citations, distinguish independent corroboration from repetition, and compare research findings with later representations.',
'Searching as Strategic Exploration — very strong: students choose what to check, trace sources laterally, use authoritative channels, revise search paths, and decide when evidence is sufficient.'
]: insert_before(anchor,text,'List Bullet')
insert_before(anchor,'The 2026 revision is especially relevant to Evidence Lab because it explicitly addresses emerging technologies and AI, metacognition, evolving information creation processes, and critical evaluation of AI-mediated searching and authority.','Normal')
insert_before(anchor,'Framework reference: Association of College and Research Libraries. Framework for Information Literacy for Higher Education. Adopted 2016; revisions approved September 2, 2026.','Normal')
doc.save(docp)

# Audit
checks={
'case18_value_language': any(c['bank_number']==18 and 'financial relationships' in c['transfer_principle'] for c in d['cases']),
'case21_value_language': any(c['bank_number']==21 and 'Money and personal information both have value' in c['transfer_principle'] for c in d['cases']),
'play_synced': 'financial relationships surrounding the recommendation' in s and 'Money and personal information both have value' in s,
'readme_alignment':'ACRL Framework' in readme.read_text(encoding='utf-8'),
'educator_page_alignment':'ACRL Framework' in edu.read_text(encoding='utf-8'),
'confidence_patch_preserved':'Choose confidence…' in s and 'Choose a confidence level before submitting.' in s,
'strategy_preserved':'Strategy ${strategy}/1' in s,
'log_preserved':'Investigation Log' in s
}
(root/'ACRL_2026_ALIGNMENT_AUDIT.json').write_text(json.dumps({'status':'PASS' if all(checks.values()) else 'BLOCKER','checks':checks},indent=2),encoding='utf-8')
print(checks)
